"""Turn uploaded files into text + structured rows.

PDF/CSV/XLSX all converge on the same output shape so the extraction
pipeline downstream doesn't care what the source format was.
"""
import csv
import io
import logging
from dataclasses import dataclass, field

log = logging.getLogger("trinetra.parsers")


@dataclass
class ParsedDocument:
    kind: str                       # fir | cdr | bank | vehicle | criminal_history | generic
    pages: list[str] = field(default_factory=list)     # text per page
    rows: list[dict] = field(default_factory=list)     # structured rows
    columns: list[str] = field(default_factory=list)
    meta: dict = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        return "\n\n".join(self.pages)


# ------------------------------------------------------------------ PDF
def parse_pdf(data: bytes) -> ParsedDocument:
    try:
        import pdfplumber
    except ImportError:
        return ParsedDocument(kind="fir", pages=[], meta={"error": "pdfplumber missing"})

    pages: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return ParsedDocument(kind="fir", pages=pages, meta={"page_count": len(pages)})


# ------------------------------------------------------------- CSV/XLSX
def _sniff_kind(columns: list[str]) -> str:
    cols = {c.lower().strip() for c in columns}
    if {"caller", "callee"} & cols or {"a_party", "b_party"} & cols:
        return "cdr"
    if {"amount", "debit", "credit"} & cols:
        return "bank"
    if {"registration_no", "vehicle_no", "plate"} & cols:
        return "vehicle"
    if {"offense_type", "offence_type", "prior_case"} & cols:
        return "criminal_history"
    return "generic"


def parse_csv(data: bytes) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    rows = [dict(r) for r in reader]
    columns = list(reader.fieldnames or [])
    return ParsedDocument(kind=_sniff_kind(columns), rows=rows,
                          columns=columns, meta={"row_count": len(rows)})


def parse_xlsx(data: bytes) -> ParsedDocument:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ParsedDocument(kind="generic", meta={"error": "openpyxl missing"})

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    try:
        header = [str(h) if h is not None else "" for h in next(it)]
    except StopIteration:
        return ParsedDocument(kind="generic")

    rows = []
    for raw in it:
        if raw is None or all(v is None for v in raw):
            continue
        rows.append({header[i]: ("" if v is None else str(v))
                     for i, v in enumerate(raw) if i < len(header)})
    return ParsedDocument(kind=_sniff_kind(header), rows=rows,
                          columns=header, meta={"row_count": len(rows)})


def parse_txt(data: bytes) -> ParsedDocument:
    return ParsedDocument(kind="generic",
                          pages=[data.decode("utf-8", errors="replace")])


# ---------------------------------------------------------------- audio
def parse_audio(data: bytes, filename: str) -> ParsedDocument:
    """Offline transcription via faster-whisper. Nothing leaves the deployment.

    Degrades honestly if the model isn't installed (free-tier RAM limits).
    """
    from app.core.config import settings
    if not settings.WHISPER_ENABLED:
        return ParsedDocument(kind="audio", pages=[],
                              meta={"skipped": "WHISPER_ENABLED=false"})
    try:
        import tempfile
        from faster_whisper import WhisperModel
    except ImportError:
        return ParsedDocument(kind="audio", pages=[],
                              meta={"error": "faster-whisper not installed"})

    with tempfile.NamedTemporaryFile(suffix=filename[-5:], delete=False) as tmp:
        tmp.write(data)
        path = tmp.name

    model = WhisperModel(settings.WHISPER_MODEL_SIZE, device="cpu",
                         compute_type="int8")
    segments, info = model.transcribe(path, beam_size=1)
    text = " ".join(s.text for s in segments)
    return ParsedDocument(kind="audio", pages=[text],
                          meta={"language": info.language,
                                "duration": info.duration,
                                "model": settings.WHISPER_MODEL_SIZE})


EXT_MAP = {
    ".pdf": parse_pdf, ".csv": parse_csv, ".xlsx": parse_xlsx,
    ".xlsm": parse_xlsx, ".txt": parse_txt, ".md": parse_txt,
}
AUDIO_EXT = {".wav", ".mp3", ".m4a", ".ogg", ".flac"}


def parse_file(filename: str, data: bytes) -> ParsedDocument:
    lower = filename.lower()
    for ext, fn in EXT_MAP.items():
        if lower.endswith(ext):
            return fn(data)
    for ext in AUDIO_EXT:
        if lower.endswith(ext):
            return parse_audio(data, filename)
    return parse_txt(data)
